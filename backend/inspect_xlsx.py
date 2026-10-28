import openpyxl

wb = openpyxl.load_workbook("uploads/invoice.xlsx", read_only=True)
sheet = wb[wb.sheetnames[0]]
cols = []
for i, row in enumerate(sheet.iter_rows(values_only=True)):
    if i == 0:
        cols = row
    elif i <= 5:
        d = {cols[j]: row[j] for j in range(min(len(cols), len(row)))}
        inv_id = d.get("IDInvoice")
        doc_no = d.get("DocumentNo")
        dt = d.get("DateInvoice")
        cust = str(d.get("NameCustomers"))[:25]
        amt = d.get("SalesAmount")
        prod = str(d.get("DescriptionProduit"))[:25]
        print(f"Row {i} -> Inv: {inv_id} | Doc: {doc_no} | Date: {dt} | Cust: {cust} | SalesAmt: {amt} | Prod: {prod}")
